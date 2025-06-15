from customtkinter import *
from json import loads, dumps
from Widgets import Entries
from Proccesors.Database.db import *
from PIL import Image
from ServicesApplier.TestManagment.Quiz import QuizApplier
from Views import Sessions, Dashboard
import ttkbootstrap as ttk
from ttkbootstrap.constants import *
import customtkinter as ctk
import tkinter as tk
from Widgets import Scrolls
import tkinter.messagebox as messagebox
from CTkTable import CTkTable
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from datetime import datetime
import os
from Views.ExcelReader import ExcelReader

def export_to_excel(exam_info, exam_newData):
    try:
        # Create a new workbook and select the active sheet
        wb = Workbook()
        ws = wb.active
        ws.title = "تقرير درجات الامتحان"
        
        # Set RTL direction for the sheet
        ws.sheet_view.rightToLeft = True

        # Define styles
        header_fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
        header_font = Font(name='Cairo', size=12, bold=True, color="FFFFFF")
        cell_font = Font(name='Cairo', size=11)
        pass_fill = PatternFill(start_color="C8E6C9", end_color="C8E6C9", fill_type="solid")
        fail_fill = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
        
        # Cell alignment
        alignment = Alignment(horizontal='center', vertical='center')
        
        # Border style
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Write exam information
        ws['A1'] = "تقرير درجات الامتحان"
        ws['A1'].font = Font(name='Cairo', size=16, bold=True)
        ws.merge_cells('A1:D1')
        ws['A1'].alignment = alignment

        ws['A2'] = f"اسم الامتحان: {exam_info[1]}"
        ws['A3'] = f"الدرجة النهائية: {exam_info[2]}"
        ws['A4'] = f"الحالة: {'مكتمل' if exam_info[4] == 'DONE' else 'مؤرشف أو لم يتم اكماله'}"
        ws['A5'] = f"تاريخ التصدير: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

        for cell in [ws['A2'], ws['A3'], ws['A4'], ws['A5']]:
            cell.font = cell_font
            cell.alignment = alignment

        # Headers
        headers = ["اسم الطالب", "الدرجة", "النسبة المئوية", "الحالة"]
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=7, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = alignment
            cell.border = border
            ws.column_dimensions[chr(64 + col)].width = 20

        # Data
        row = 8
        sorted_degrees = sorted(exam_newData['degrees'], key=lambda d: int(list(d.values())[0]), reverse=True)
        for degree in sorted_degrees:
            student_name = list(degree.keys())[0]
            student_grade = int(list(degree.values())[0])
            percentage = (student_grade / int(exam_newData['max'])) * 100
            status = "ناجح" if student_grade > int(exam_newData['max'])/2 else "راسب"
            
            ws.cell(row=row, column=1, value=student_name)
            ws.cell(row=row, column=2, value=student_grade)
            ws.cell(row=row, column=3, value=f"{percentage:.1f}%")
            ws.cell(row=row, column=4, value=status)
            
            # Apply styles
            fill = pass_fill if status == "ناجح" else fail_fill
            for col in range(1, 5):
                cell = ws.cell(row=row, column=col)
                cell.fill = fill
                cell.font = cell_font
                cell.alignment = alignment
                cell.border = border
            
            row += 1

        # Add summary statistics
        total_students = len(exam_newData['degrees'])
        passing_students = sum(1 for deg in exam_newData['degrees'] 
                             if int(list(deg.values())[0]) > int(exam_newData['max'])/2)
        
        summary_row = row + 2
        ws.cell(row=summary_row, column=1, value="إحصائيات عامة")
        ws.merge_cells(f'A{summary_row}:D{summary_row}')
        ws.cell(row=summary_row, column=1).font = Font(name='Cairo', size=14, bold=True)
        ws.cell(row=summary_row, column=1).alignment = alignment

        stats = [
            ("عدد الطلاب الكلي", total_students),
            ("عدد الطلاب الناجحين", passing_students),
            ("عدد الطلاب الراسبين", total_students - passing_students),
            ("نسبة النجاح", f"{(passing_students/total_students*100):.1f}%" if total_students > 0 else "0%")
        ]

        for i, (label, value) in enumerate(stats):
            ws.cell(row=summary_row + i + 1, column=1, value=label)
            ws.cell(row=summary_row + i + 1, column=2, value=str(value))
            ws.merge_cells(f'B{summary_row + i + 1}:D{summary_row + i + 1}')
            ws.cell(row=summary_row + i + 1, column=1).font = cell_font
            ws.cell(row=summary_row + i + 1, column=2).font = cell_font
            ws.cell(row=summary_row + i + 1, column=1).alignment = alignment
            ws.cell(row=summary_row + i + 1, column=2).alignment = alignment

        # Save the file
        if not os.path.exists('Reports'):
            os.makedirs('Reports')
            
        filename = f"Reports/تقرير_درجات_{exam_info[1]}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(filename)
        messagebox.showinfo("تم", f"تم تصدير التقرير بنجاح إلى:\n{filename}")
        # Open the saved file in a topmost ExcelReader
        top = CTkToplevel()
        top.title(f"عرض ملف Excel: {filename}")
        top.attributes("-topmost", True)
        top.geometry("900x600")
        frame = CTkScrollableFrame(top, fg_color="white")
        frame.pack(fill="both", expand=True, padx=10, pady=10)
        ExcelReader(frame, file_path=filename)
        
    except Exception as e:
        messagebox.showerror("خطأ", f"حدث خطأ أثناء تصدير التقرير: {str(e)}")

def show_exam_info(frame : CTkScrollableFrame, exam_info):
        for x in frame.winfo_children():
            x.destroy()

        try:
            ref.execute("PRAGMA database_list;")
            db_info = ref.fetchone()
            if not db_info:
                messagebox.showerror("خطأ", "لا يمكن الوصول إلى قاعدة البيانات")
                return
                
            team = ((db_info[2]).split("\\")[-1]).replace(".db", "")
            conn_checker = connect("Application.db")
            ref_checker = conn_checker.cursor()
            ref_checker.execute("SELECT * FROM PGroups WHERE Name = ?", (team,))
            
            group_info = ref_checker.fetchone()
            if not group_info:
                type_of_payment = "غير محدد"  # Default value if not found
            else:
                type_of_payment = group_info[2]

            ref.execute('SELECT * FROM Sessions WHERE ID = ?', (loads(exam_info[5])['id'],))
            session_data = ref.fetchone()
            if not session_data:
                messagebox.showerror("خطأ", "لا يمكن العثور على بيانات الجلسة")
                return
                
            exam_newData = loads(session_data[7])
            
            # Create main info frame
            info_frame = CTkFrame(frame, fg_color="white")
            info_frame.pack(fill=X, padx=10, pady=10)
            
            # Add export button
            export_button = CTkButton(
                info_frame,
                text="تصدير تقرير Excel",
                font=('Cairo Medium', 16),
                fg_color="#4CAF50",
                hover_color="#388E3C",
                command=lambda: export_to_excel(exam_info, exam_newData)
            )
            export_button.pack(side=TOP, pady=(10,5))
            
            CTkLabel(info_frame, text=exam_info[1], font=('FF Shamel Family Sans One Bold', 40)).pack(pady=(20,10))
            CTkLabel(info_frame, text=f"الدرجات المرصودة : {len(exam_newData['degrees'])}", font=('FF Shamel Family Sans One Book', 17)).pack()
            CTkLabel(info_frame, text=f"الحالة : {'مكتمل' if exam_info[4] == 'DONE' else 'مؤرشف أو لم يتم اكماله'}", font=('FF Shamel Family Sans One Book', 17), text_color=('green' if exam_info[4] == 'DONE' else 'ORANGE')).pack()
            CTkLabel(info_frame, text=f"الدرجة النهائية : {exam_info[2]}", font=('FF Shamel Family Sans One Book', 17)).pack()
            
            success_rate = 0
            for deg in exam_newData['degrees']:
                if int(list(deg.values())[0]) > int(exam_newData['max'])/2:
                    success_rate += 1
            try:
                success_rate = success_rate/int(len(exam_newData['degrees']))
            except ZeroDivisionError:
                success_rate = 0

            # Statistics Frame
            percentage = CTkFrame(frame, fg_color='white')
            stctsfrm = CTkFrame(percentage, fg_color='transparent')
            stctsfrm.pack()
            
            # Success Rate Circle
            yattend = CTkFrame(stctsfrm, fg_color='transparent')
            yattend.pack(fill=BOTH, side=RIGHT,pady=(10,10),padx=(0,5))
            circular = CTkFrame(yattend, fg_color='#36BA98', corner_radius=180, width=100, height=100)
            circular.pack(pady=(15,10))
            attendplbl = CTkLabel(circular, text=f'{int(success_rate*100)}%', font=('Arial', 45, 'bold'), text_color='white')
            attendplbl.pack(fill=BOTH, expand=True, padx=25,pady=43)
            CTkLabel(yattend, text='نسبة النجاح', font=('Cairo Medium', 20)).pack()

            # Failure Rate Circle
            yattend = CTkFrame(stctsfrm, fg_color='transparent')
            yattend.pack(fill=BOTH, side=RIGHT,pady=(10,10),padx=(5,20))
            circular = CTkFrame(yattend, fg_color='orange', corner_radius=180, width=100, height=100)
            circular.pack(pady=(15,10))
            attendplbl = CTkLabel(circular, text=f'{100 - int(success_rate*100)}%', font=('Arial', 45, 'bold'), text_color='white')
            attendplbl.pack(fill=BOTH, expand=True, padx=25,pady=43)
            CTkLabel(yattend, text='نسبة الرسوب', font=('Cairo Medium', 20)).pack()

            percentage.pack(pady=(20,5), fill=X)

            # Add Student Grades Table
            table_frame = CTkFrame(frame, fg_color='white')
            table_frame.pack(fill=BOTH, expand=True, padx=10, pady=10)
            
            CTkLabel(table_frame, text="تفاصيل درجات الطلاب", font=('Cairo Medium', 24)).pack(pady=(10,20))
            
            # Create table headers
            headers = ["اسم الطالب", "الدرجة", "الحالة"]
            table_data = [headers]
            
            # Prepare table data
            sorted_degrees = sorted(exam_newData['degrees'], key=lambda d: int(list(d.values())[0]), reverse=True)
            for degree in sorted_degrees:
                student_name = list(degree.keys())[0]
                student_grade = list(degree.values())[0]
                status = "ناجح" if int(student_grade) > int(exam_newData['max'])/2 else "راسب"
                status_color = "green" if status == "ناجح" else "red"
                table_data.append([student_name, student_grade, status])

            # Create table
            grades_table = CTkTable(
                table_frame,
                values=table_data,
                colors=["#E3E3E3", "#EEEEEE"],
                header_color="#D1D1D1",
                hover_color="#CCCCCC",
                font=("Cairo Medium", 16),
                width=120,
                height=35,
                corner_radius=0
            )
            grades_table.pack(fill=BOTH, expand=True, padx=10, pady=10)

            if exam_info[4] == 'DONE':
                CTkButton(frame, text='رصد او رؤية الدرجات',command=lambda:QuizApplier(session_data).AddDegrees(), font=('Cairo Medium', 20)).pack(fill=X, padx=40, ipady=7, pady=20)
            else:
                CTkButton(frame, text='رصد او رؤية الدرجات',command=lambda:QuizApplier(session_data).AddDegrees(), font=('Cairo Medium', 20)).pack(fill=X, padx=40, ipady=7, pady=20)

        except Exception as e:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء عرض معلومات الامتحان: {str(e)}")
        finally:
            if 'conn_checker' in locals():
                conn_checker.close()

def go(frame : CTkScrollableFrame):
    for widget in frame.winfo_children():
        widget.destroy()
    ss = CTkFrame(frame, fg_color='white')
    sessions_listview = CTkScrollableFrame(ss, height=600)
    statics = CTkFrame(frame, fg_color='white')
    ref.execute('SELECT * FROM Exams')
    sessions = ref.fetchall()
#, text='تحديد الامتحان كمكتمـل', font=('Cairo Medium', 20), fg_color='#399918', hover_color='#1A5319').pack(fill=X, padx=40, ipady=7)
    if sessions == []:
        CTkLabel(sessions_listview, text='', image=CTkImage(Image.open('TDSAssets/General/calender.png'), size=(120,120))).pack(pady=(150,5))
        CTkLabel(sessions_listview, text= 'لا توجد امتحانات هنا', font=('Cairo Medium',23)).pack()
    else:
        for x in sessions_listview.winfo_children():
            x.destroy()
        ref.execute(f'SELECT * FROM Exams')
        sessions = ref.fetchall()
        for session in sessions:
            ff = CTkFrame(sessions_listview, fg_color='white')
            ff.pack(fill=X, padx=5, pady=5)
            f1 = CTkFrame(ff, fg_color='transparent')
            f1.pack(fill=X, pady=2)
            f2 = CTkFrame(ff, fg_color='transparent')
            f2.pack(fill=X, pady=2)
            CTkLabel(f1, text=session[1], font=('Cairo Medium', 20)).pack(pady=5, padx=15, side=RIGHT)
            price_status = "لم يتم رصد كل الدرجات , مسودة \ ارشيف" if session[4] == "PENDING" else "مكتمل"
            price_color = "orange" if session[4] == "PENDING" else "green"
            CTkLabel(f2, text=price_status, font=('Cairo Medium', 13), text_color=price_color).pack(pady=5, padx=15, side=LEFT)

            if session[5] != "":
                session_data = loads(session[5])
                ref.execute(f"SELECT * FROM Sessions WHERE ID='{session_data['id']}'")
                related_session = ref.fetchone()
                if related_session:
                    CTkLabel(f2, text=f"للحصة : {related_session[1]}", font=('Cairo Medium', 13)).pack(pady=5, padx=15, side=RIGHT)

            CTkButton(f1, text='تفاصيل', font=('Cairo Medium', 18), command=lambda ex = session:show_exam_info(statics, ex)).pack(pady=5, padx=15, side=LEFT)
            
    kew = Entries.UpperLabeledEntry(ss, 'بحث عن امتحان')
    kew.pack(fill=BOTH,padx=10, pady=(10,0))
    def SearchSession(event):
        if event.widget.get() == '':
            for x in sessions_listview.winfo_children():
                x.destroy()
            ref.execute('SELECT * FROM Exams')
            sessions = ref.fetchall()
            if not sessions:
                CTkLabel(sessions_listview, text='', image=CTkImage(Image.open('TDSAssets/General/calender.png'), size=(120,120))).pack(pady=(150,5))
                CTkLabel(sessions_listview, text='لا توجد امتحانات هنا', font=('Cairo Medium', 23)).pack()
            else:
                for x in sessions_listview.winfo_children():
                    x.destroy()
                ref.execute(f'SELECT * FROM Exams WHERE Name LIKE "%{event.widget.get()}%"')
                sessions = ref.fetchall()
                for session in sessions:
                    ff = CTkFrame(sessions_listview, fg_color='white')
                    ff.pack(fill=X, padx=5, pady=5)
                    f1 = CTkFrame(ff, fg_color='transparent')
                    f1.pack(fill=X, pady=2)
                    f2 = CTkFrame(ff, fg_color='transparent')
                    f2.pack(fill=X, pady=2)
                    CTkLabel(f1, text=session[1], font=('Cairo Medium', 20)).pack(pady=5, padx=15, side=RIGHT)
                    price_status = "لم يتم رصد كل الدرجات , مسودة \ ارشيف" if session[4] == "PENDING" else "مكتمل"
                    price_color = "orange" if session[4] == "PENDING" else "green"
                    CTkLabel(f2, text=price_status, font=('Cairo Medium', 13), text_color=price_color).pack(pady=5, padx=15, side=LEFT)

                    if session[5] != "":
                        session_data = loads(session[5])
                        ref.execute(f"SELECT * FROM Sessions WHERE ID='{session_data['id']}'")
                        related_session = ref.fetchone()
                        if related_session:
                            CTkLabel(f2, text=f"للحصة : {related_session[1]}", font=('Cairo Medium', 13)).pack(pady=5, padx=15, side=RIGHT)

                    CTkButton(f1, text='تفاصيل', font=('Cairo Medium', 18), command=lambda ex = session:show_exam_info(statics, ex)).pack(pady=5, padx=15, side=LEFT)
                    
        else:
            for x in sessions_listview.winfo_children():
                x.destroy()
            ref.execute(f'SELECT * FROM Exams WHERE Name LIKE "%{event.widget.get()}%"')
            sessions = ref.fetchall()
            for session in sessions:
                ff = CTkFrame(sessions_listview, fg_color='white')
                ff.pack(fill=X, padx=5, pady=5)
                f1 = CTkFrame(ff, fg_color='transparent')
                f1.pack(fill=X, pady=2)
                f2 = CTkFrame(ff, fg_color='transparent')
                f2.pack(fill=X, pady=2)
                CTkLabel(f1, text=session[1], font=('Cairo Medium', 20)).pack(pady=5, padx=15, side=RIGHT)
                price_status = "لم يتم رصد كل الدرجات , مسودة \ ارشيف" if session[4] == "PENDING" else "مكتمل"
                price_color = "orange" if session[4] == "PENDING" else "green"
                CTkLabel(f2, text=price_status, font=('Cairo Medium', 13), text_color=price_color).pack(pady=5, padx=15, side=LEFT)

                if session[5] != "":
                    session_data = loads(session[5])
                    ref.execute(f"SELECT * FROM Sessions WHERE ID='{session_data['id']}'")
                    related_session = ref.fetchone()
                    if related_session:
                        CTkLabel(f2, text=f"للحصة : {related_session[1]}", font=('Cairo Medium', 13)).pack(pady=5, padx=15, side=RIGHT)

                CTkButton(f1, text='تفاصيل', font=('Cairo Medium', 18), command=lambda ex = session:show_exam_info(statics, ex)).pack(pady=5, padx=15, side=LEFT)
                
            if len(sessions_listview.winfo_children()) == 0:
                CTkLabel(sessions_listview, text='', image=CTkImage(Image.open('TDSAssets/General/error.png'), size=(90,90))).pack(pady=(60,5))
                CTkLabel(sessions_listview, text='هذا الامتحان غير موجود', font=('Cairo Medium', 19)).pack()

    kew.entry.bind('<KeyRelease>', SearchSession)
    sessions_listview.pack(pady=10, fill=BOTH, expand=True, padx=20)
    ss.grid(row=0, column=1, sticky='nsew', padx=10, pady=10)



    statics.grid(row=0, column=0, sticky='nsew', padx=10, pady=10)
    CTkLabel(statics, text='', image=CTkImage(Image.open('TDSAssets/Banking/1.png'), size=(404, 580))).pack(fill=BOTH, expand=True, pady=5,padx=5)
    def back():
        for x in frame.winfo_children():
            x.destroy()
        Dashboard.Dashboard(frame).pack(fill=BOTH, expand=True)
    CTkButton(ss, text='الرجوع الي الصفحة الرئيسية   ', font=('Cairo Medium', 16),text_color='black', fg_color='#80C4E9', hover=False, image=CTkImage(Image.open('TDSAssets/General/backbutton.png'), size=(40,40)), compound='right', command=lambda:back()).pack(fill=X, padx=5, pady=5, ipady=5, side=BOTTOM)

    frame.columnconfigure(0, weight=1)
    frame.columnconfigure(1, weight=1)
    frame.rowconfigure(0, weight=1)
    frame.rowconfigure(1, weight=2)





    
