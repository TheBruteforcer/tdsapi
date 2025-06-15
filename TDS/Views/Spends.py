from customtkinter import *
from PIL import Image
from Views import Dashboard
from time import localtime
from datetime import datetime, timedelta
from Proccesors.Database.db import ref
from json import loads
from datetime import timedelta
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment

def show(frame: CTkFrame):
    def back():
        for x in frame.winfo_children():
            x.destroy()
        Dashboard.Dashboard(frame).pack(fill=BOTH, expand=True)

    def get_date_range(key):
        today = datetime.today()
        if key == 'يوم':
            start_date = end_date = today
        elif key == 'اسبوع':
            end_date = today - timedelta(days=today.weekday() + 1)
            start_date = end_date - timedelta(days=6)
        elif key == 'شهر':
            start_date = today.replace(day=1)
            next_month = today.replace(day=28) + timedelta(days=4)
            end_date = next_month - timedelta(days=next_month.day)
        elif key == 'سنة':
            start_date = today.replace(month=1, day=1)
            end_date = today.replace(month=12, day=31)
        else:
            start_date = end_date = today  # Default to today if key is not recognized
        return start_date, end_date

    def format_date(date):
        return f'{date.day} \ {date.month} \ {date.year}'

    def filter_data(key):
        key = key.strip()
        start_date, end_date = get_date_range(key)
        total_spends = 0
        for x in spendslist.winfo_children():
            x.destroy()
        for z in total.winfo_children():
            z.destroy()
        
        current_date = start_date
        spends = []
        while current_date <= end_date:
            print(format_date(current_date))
            formatted_date = format_date(current_date)
            ref.execute(f'SELECT * FROM Spends WHERE Date = "{formatted_date}"')
            spends.extend(ref.fetchall())
            current_date += timedelta(days=1)

        if spends:
            for x in spends:
                total_spends += int(x[0])
                ff = CTkFrame(spendslist, fg_color='white')
                ff.pack(padx=5, pady=5, fill=X)
                CTkLabel(ff, text='', image=CTkImage(Image.open('TDSAssets/General/red_down.png'), size=(40, 40))).pack(pady=5, padx=5, side=RIGHT)
                CTkLabel(ff, text=loads(x[2])['inf'], font=('Cairo Medium', 20)).pack(side=RIGHT, padx=10)
                CTkLabel(ff, text=loads(x[2])['nm'], font=('Cairo Medium', 20), text_color='green').pack(side=RIGHT, padx=10)
                CTkLabel(ff, text=f'E£ {int(x[0]):,}', font=('Cairo Medium', 18)).pack(side=RIGHT, padx=10)
                CTkLabel(ff, text=f'الوقت : {x[1]}', font=('Cairo Medium', 18)).pack(side=RIGHT, padx=10)
            CTkLabel(total, text=f'المجموع : {total_spends} جنيه مصري', font=('Cairo Medium', 18)).pack(side=RIGHT, pady=5, padx=20)
            CTkButton(total, text="استخراج تقرير", font=("Cairo Medium", 16), command = export).pack(side=LEFT)

        else:
            CTkLabel(total, text='غير متوفر', font=('Cairo Medium', 18)).pack(side=RIGHT, pady=5, padx=20)
            ff = CTkFrame(spendslist, fg_color='transparent')
            ff.pack(pady=120)
            CTkLabel(ff, text='', image=CTkImage(Image.open('TDSAssets/General/red_bag.png'), size=(40, 40))).pack(pady=5, padx=5, side=RIGHT)
            CTkLabel(ff, text='ستظهر المصروفات هنا عند الإضافة', font=('Cairo Medium', 20)).pack(side=RIGHT, padx=10)

    for x in frame.winfo_children():
        x.destroy()
    
    title = CTkFrame(frame)
    title.pack(side=TOP, fill=X, padx=10, pady=10)
    CTkLabel(title, text='', image=CTkImage(Image.open('TDSAssets/General/red_bag.png'), size=(60, 60))).pack(side=RIGHT, pady=10, padx=(15, 10))
    tt = CTkFrame(title, fg_color='transparent')
    tt.pack(side=RIGHT)
    CTkLabel(tt, text='المصروفات', font=('Cairo Medium', 15), compound='right', anchor='e').pack(fill=X)
    
    today = localtime()
    formatted_today = f'{today.tm_mday} \ {today.tm_mon} \ {today.tm_year}'
    total_spends = 0
    ref.execute(f'SELECT * FROM Spends WHERE Date = "{formatted_today}"')
    for x in ref.fetchall():
        total_spends += int(x[0])
    CTkLabel(tt, text=f'المبلغ المصروف اليوم : {total_spends}', font=('Cairo Medium', 14), compound='right', anchor='e').pack(fill=X)
    
    CTkButton(title, text='العودة الي الصفحة الرئيسية', font=('Cairo Medium', 17), command=back).pack(side=LEFT, padx=15)
    
    listingfrm = CTkFrame(frame, fg_color='transparent')
    listingfrm.pack(fill=X, pady=10, padx=10)
    
    filterfrm = CTkFrame(listingfrm)
    filterfrm.pack(fill=X)
    filterI = StringVar(value='  يوم  ')
    spendslist = CTkScrollableFrame(listingfrm, height=430, label_fg_color='#F4CE14', label_text='المصروفات', label_font=('Cairo Medium', 15))
    
    CTkLabel(filterfrm, text='فلترة في مدة زمنية', font=('Cairo Medium', 18)).pack(side=RIGHT, pady=5, padx=(0, 10))
    CTkOptionMenu(filterfrm, command=filter_data, variable=filterI, values=['  يوم  ', '  اسبوع  ', '  شهر  ', '  سنة  '], font=('Cairo Medium', 16), anchor='e').pack(side=LEFT, padx=10)
    
    total = CTkFrame(listingfrm)
    filter_data('يوم')
    spendslist.pack(pady=10, fill=X)


    def export():
        key = filterI.get().strip()  # Get selected filter option
        start_date, end_date = get_date_range(key)  # Get date range
        result_list = []  # Store results

        current_date = start_date
        while current_date <= end_date:
            formatted_date = format_date(current_date)
            ref.execute(f'SELECT * FROM Spends WHERE Date = "{formatted_date}"')
            result_list.extend(ref.fetchall())  # Append results
            current_date += timedelta(days=1)

        formatted_results = []  # Store formatted data
        for spend in result_list:
            spend_amount  = spend[0]
            spend_date    = spend[1]
            spender_name  = ''
            spend_desc    = ''
            spend_session = ''
            spend_group   = ''
            note          = ''

            try:
                spender_name  = loads(spend[2])["nm"]
                spend_desc    = loads(spend[2])["inf"]
                spend_session = loads(spend[3])["session-info"]["title"]
                ref.execute("SELECT * FROM Sessions WHERE ID = ?", (loads(spend[3])["session-info"]["id"],))
                spend_group   = ref.fetchone()[10]
            except:
                note = "استخراج مستلزمات"

            formatted_results.append([spend_date, spend_amount, spend_desc, spender_name, spend_session, spend_group, note])

        # Load the workbook and select the active sheet
        file_path = "Templates/DetailedSpendsAccounting.xlsx"
        wb = load_workbook(file_path)
        ws = wb.active  # Change to a specific sheet with wb["SheetName"] if needed

        # Start inserting data from row 5
        start_row = 5  
        for row_index, row_data in enumerate(formatted_results, start=start_row):
            for col_index, cell_value in enumerate(row_data, start=1):  
                cell = ws.cell(row=row_index, column=col_index, value=cell_value)
                cell.font = Font(name="Tajawal", size=12)  # Set font
                cell.alignment = Alignment(horizontal="center", vertical="center")  # Center text

            ws.row_dimensions[row_index].height = 20  # Adjust row height

        # Save the updated workbook
        output_path = "التقارير/تقرير_مصروفات.xlsx"
        wb.save(output_path)

        print(f"تم حفظ التقرير بنجاح في: {output_path}")
        from Views.ExcelReader import ExcelReader
        top = CTkToplevel()
        top.title(f"عرض ملف Excel: {output_path}")
        top.attributes("-topmost", True)
        top.geometry("900x600")
        frame = CTkScrollableFrame(top, fg_color="white")
        frame.pack(fill="both", expand=True, padx=10, pady=10)
        ExcelReader(frame, file_path=output_path)

    CTkButton(total, text="استخراج تقرير", font=("Cairo Medium", 16), command = export).pack(side=LEFT)
    total.pack(fill=X)
