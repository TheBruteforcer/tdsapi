from customtkinter import *
from PIL import Image
from Widgets import Entries
from Proccesors import *
from requests import *
from Proccesors.Database import db
from threading import Thread
from Proccesors.Genric.SearchStudent import *
from AlphaControllers.EntryValidators import CheckBlank
from Views import Dashboard,Admin,AddStdForm, Spends, AddSpend, Sessions, SessionsSettings, Exams, Money, Groups, AllStudents, Teachers
from time import localtime
from ServicesApplier.MLPowerdSelection import View
from Widgets import Entries
from time import localtime
from AlphaControllers import EntryValidators
import os
from datetime import datetime
from sqlite3 import connect
from tkinter import messagebox
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from Views.PendingCodes import PendingCodesView
from Views.ExcelReader import ExcelReader

class TeamView(CTkFrame):
    def __init__(self, parent : CTkFrame):
        self.ss = parent
        super().__init__(parent, fg_color='#F0EBE3')
        
        # Header Frame
        header_frame = CTkFrame(self, fg_color="transparent")
        header_frame.pack(fill=X, pady=(20,0))
        
        # Title
        CTkLabel(header_frame, text="إختر مرحلة", font=("Cairo Medium", 60)).pack(side=RIGHT, pady=(10,0), padx=20)
        
        # Admin Button with modern styling
        admin_icon = CTkImage(Image.open("TDSAssets/General/settings.png"), size=(25, 25))
        admin_btn = CTkButton(
            header_frame,
            text="صفحة المدير",
            font=("Cairo Medium", 18),
            image=admin_icon,
            compound="right",
            command=self.open_admin,
            fg_color="#2B6BE6",
            hover_color="#1E88E5",
            corner_radius=12,
            height=45
        )
        admin_btn.pack(side=LEFT, pady=(10,0), padx=20)
        
        # Custom Gains Button
        gains_icon = CTkImage(Image.open("TDSAssets/General/green_money.png"), size=(25, 25))
        CTkButton(
            header_frame,
            text="الدخل الإضافي",
            font=("Cairo Medium", 18),
            image=gains_icon,
            compound="right",
            command=self.open_custom_gains,
            fg_color="#4CAF50",
            hover_color="#45a049",
            corner_radius=12,
            height=45
        ).pack(side=LEFT, pady=(10,0), padx=5)

        # Pending Codes Button
        codes_icon = CTkImage(Image.open("TDSAssets/General/note.png"), size=(25, 25))
        CTkButton(
            header_frame,
            text="الكروت الغير مطبوعة",
            font=("Cairo Medium", 18),
            image=codes_icon,
            compound="right",
            command=self.open_pending_codes,
            fg_color="#FF9800",
            hover_color="#FFA726",
            corner_radius=12,
            height=45
        ).pack(side=LEFT, pady=(10,0), padx=5)

        # Stats Card
        self.stats_frame = CTkFrame(self, fg_color="#FFFFFF", corner_radius=15)
        self.stats_frame.pack(fill=X, padx=20, pady=10)
        
        # Create three columns for stats
        gains_frame = CTkFrame(self.stats_frame, fg_color="transparent")
        gains_frame.pack(side=LEFT, expand=True, padx=20, pady=10)
        
        sessions_frame = CTkFrame(self.stats_frame, fg_color="transparent")
        sessions_frame.pack(side=LEFT, expand=True, padx=20, pady=10)
        
        students_frame = CTkFrame(self.stats_frame, fg_color="transparent")
        students_frame.pack(side=LEFT, expand=True, padx=20, pady=10)
        
        # Stats Labels
        self.gains_label = CTkLabel(gains_frame, text="0", font=("Cairo Bold", 24))
        self.gains_label.pack()
        CTkLabel(gains_frame, text="أرباح اليوم", font=("Cairo Medium", 16)).pack()
        CTkButton(
            gains_frame,
            text="تصدير البيانات",
            font=("Cairo Medium", 14),
            command=self.export_gains,
            fg_color="#4CAF50",
            hover_color="#45a049",
            corner_radius=8,
            height=30
        ).pack(pady=(10,0))
        
        self.sessions_label = CTkLabel(sessions_frame, text="0", font=("Cairo Bold", 24))
        self.sessions_label.pack()
        CTkLabel(sessions_frame, text="جلسات اليوم", font=("Cairo Medium", 16)).pack()
        CTkButton(
            sessions_frame,
            text="تصدير البيانات",
            font=("Cairo Medium", 14),
            command=self.export_sessions,
            fg_color="#4CAF50",
            hover_color="#45a049",
            corner_radius=8,
            height=30
        ).pack(pady=(10,0))
        
        self.students_label = CTkLabel(students_frame, text="0", font=("Cairo Bold", 24))
        self.students_label.pack()
        CTkLabel(students_frame, text="إجمالي الطلاب", font=("Cairo Medium", 16)).pack()
        CTkButton(
            students_frame,
            text="تصدير البيانات",
            font=("Cairo Medium", 14),
            command=self.export_students,
            fg_color="#4CAF50",
            hover_color="#45a049",
            corner_radius=8,
            height=30
        ).pack(pady=(10,0))
        
        self.Buttons = CTkFrame(self, fg_color="transparent")
        self.Buttons.pack(fill=X,pady=30)
        
        CTkButton(
            self,
            command=self.AddNewTeam,
            text="إضافة مرحلة جديدة +",
            font=("Tajawal Medium", 40)
        ).pack(pady=20, ipadx=20)
        
        self.Reload()
        self.update_stats()

    def style_worksheet(self, ws, title):
        """Apply enhanced styling to worksheet"""
        # Set RTL
        ws.sheet_view.rightToLeft = True
        
        # Add title row
        ws.merge_cells('A1:E1')
        title_cell = ws['A1']
        title_cell.value = title
        title_cell.font = Font(bold=True, size=16, name='Cairo')
        title_cell.alignment = Alignment(horizontal='center', vertical='center')
        title_cell.fill = PatternFill(start_color="4CAF50", end_color="4CAF50", fill_type="solid")
        title_cell.font = Font(bold=True, size=16, name='Cairo', color="FFFFFF")
        
        # Style headers
        header_fill = PatternFill(start_color="E8F5E9", end_color="E8F5E9", fill_type="solid")
        header_font = Font(bold=True, size=12, name='Cairo')
        header_border = Border(
            bottom=Side(style='medium', color="4CAF50"),
            right=Side(style='thin', color="4CAF50"),
            left=Side(style='thin', color="4CAF50")
        )
        
        for row in ws.iter_rows(min_row=2, max_row=2):
            for cell in row:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
                cell.border = header_border
        
        # Style data cells
        data_border = Border(
            right=Side(style='thin', color="4CAF50"),
            left=Side(style='thin', color="4CAF50")
        )
        data_font = Font(name='Cairo', size=11)
        
        for row in ws.iter_rows(min_row=3):
            for cell in row:
                cell.border = data_border
                cell.font = data_font
                cell.alignment = Alignment(horizontal='center', vertical='center')
        
        # Add bottom border to last row
        for cell in ws[ws.max_row]:
            cell.border = Border(
                bottom=Side(style='medium', color="4CAF50"),
                right=Side(style='thin', color="4CAF50"),
                left=Side(style='thin', color="4CAF50")
            )
        
        # Set row height
        ws.row_dimensions[1].height = 30  # Title row
        ws.row_dimensions[2].height = 25  # Header row
        
        # Freeze header row
        ws.freeze_panes = 'A3'

    def export_gains(self):
        try:
            today = datetime.now()
            today_str = f"{today.day:02d} \\ {today.month:02d} \\ {today.year}"
            wb = Workbook()
            ws = wb.active
            ws.title = "أرباح اليوم"
            # Add title row
            ws.append([''])  # Empty row for title
            # Write headers (removed description/type column)
            headers = ['المرحلة', 'المعرف', 'القيمة', 'التاريخ']
            ws.append(headers)
            db_files = [f for f in os.listdir() if f.endswith('.db') and f != 'Application.db']
            row_count = 2  # Start after headers
            for db_file in db_files:
                try:
                    conn = connect(db_file)
                    cursor = conn.cursor()
                    # Use LIKE for date filtering
                    cursor.execute("SELECT * FROM Gains WHERE Date LIKE ?", (today_str + '%',))
                    gains = cursor.fetchall()
                    for gain in gains:
                        row_count += 1
                        ws.append([
                            db_file.replace('.db', ''),
                            gain[0],
                            f"E£ {float(gain[2]):,.2f}",
                            gain[3]  # full date+time
                        ])
                    conn.close()
                except Exception as e:
                    print(f"Error processing {db_file}: {str(e)}")
                    continue
            if row_count > 2:
                title = f"تقرير الأرباح ليوم {today_str}"
                self.style_worksheet(ws, title)
                # Adjust column widths
                for idx, column in enumerate(ws.columns, 1):
                    max_length = 0
                    column = list(column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = (max_length + 2) * 1.2  # Add some padding
                    ws.column_dimensions[get_column_letter(idx)].width = adjusted_width
                os.makedirs('exports', exist_ok=True)
                export_path = f"exports/gains_{today.strftime('%Y%m%d')}.xlsx"
                wb.save(export_path)
                messagebox.showinfo("تم", f"تم تصدير بيانات الأرباح إلى {export_path}")
                top = CTkToplevel()
                top.title(f"عرض ملف Excel: {export_path}")
                top.attributes("-topmost", True)
                top.geometry("900x600")
                frame = CTkScrollableFrame(top, fg_color="white")
                frame.pack(fill="both", expand=True, padx=10, pady=10)
                ExcelReader(frame, file_path=export_path)
            else:
                messagebox.showinfo("تنبيه", "لا توجد أرباح لتصديرها اليوم")
        except Exception as e:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء تصدير البيانات: {str(e)}")

    def export_sessions(self):
        try:
            today = datetime.now()
            today_str = f"{today.day:02d} \\ {today.month:02d} \\ {today.year}"
            
            wb = Workbook()
            ws = wb.active
            ws.title = "جلسات اليوم"
            
            # Add title row
            ws.append([''])  # Empty row for title
            
            # Write headers
            headers = ['المرحلة', 'المعرف', 'العنوان', 'المدة', 'السعر', 'التاريخ', 'وقت البدء']
            ws.append(headers)
            
            db_files = [f for f in os.listdir() if f.endswith('.db') and f != 'Application.db']
            row_count = 2  # Start after headers
            
            for db_file in db_files:
                try:
                    conn = connect(db_file)
                    cursor = conn.cursor()
                    
                    cursor.execute("""
                        SELECT ID, Title, Duration, Price, Date, StartTime 
                        FROM Sessions WHERE Date = ?
                    """, (today_str,))
                    sessions = cursor.fetchall()
                    
                    for session in sessions:
                        row_count += 1
                        ws.append([
                            db_file.replace('.db', ''),
                            session[0],
                            session[1],
                            session[2],
                            f"E£ {float(session[3]):,.2f}" if session[3] else "0",
                            session[4],
                            session[5]
                        ])
                    
                    conn.close()
                except Exception as e:
                    print(f"Error processing {db_file}: {str(e)}")
                    continue
            
            if row_count > 2:
                title = f"تقرير الجلسات ليوم {today_str}"
                self.style_worksheet(ws, title)
                
                # Adjust column widths
                for idx, column in enumerate(ws.columns, 1):
                    max_length = 0
                    column = list(column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = (max_length + 2) * 1.2
                    ws.column_dimensions[get_column_letter(idx)].width = adjusted_width
                
                os.makedirs('exports', exist_ok=True)
                export_path = f"exports/sessions_{today.strftime('%Y%m%d')}.xlsx"
                wb.save(export_path)
                messagebox.showinfo("تم", f"تم تصدير بيانات الجلسات إلى {export_path}")
                top = CTkToplevel()
                top.title(f"عرض ملف Excel: {export_path}")
                top.attributes("-topmost", True)
                top.geometry("900x600")
                frame = CTkScrollableFrame(top, fg_color="white")
                frame.pack(fill="both", expand=True, padx=10, pady=10)
                ExcelReader(frame, file_path=export_path)
            else:
                messagebox.showinfo("تنبيه", "لا توجد جلسات لتصديرها اليوم")
                
        except Exception as e:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء تصدير البيانات: {str(e)}")

    def export_students(self):
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "بيانات الطلاب"
            
            # Add title row
            ws.append([''])  # Empty row for title
            
            # Write headers
            headers = ['المرحلة', 'الكود', 'الاسم', 'رقم الهاتف', 'هاتف ولي الأمر 1', 
                       'قيمة الاشتراك']
            ws.append(headers)
            
            db_files = [f for f in os.listdir() if f.endswith('.db') and f != 'Application.db']
            row_count = 2  # Start after headers
            
            for db_file in db_files:
                try:
                    conn = connect(db_file)
                    cursor = conn.cursor()
                    
                    cursor.execute("""
                        SELECT ID, Name, OwnPhone, ParentPhone1, ParentPhone2, 
                               PayHistory, SubscribtionAmount 
                        FROM Students
                    """)
                    students = cursor.fetchall()
                    
                    for student in students:
                        row_count += 1
                        ws.append([
                            db_file.replace('.db', ''),
                            student[0],
                            student[1],
                            student[2],
                            student[3],
                            
                            f"E£ {float(student[6]):,.2f}" if student[6] else "0",
                        ])
                    
                    conn.close()
                except Exception as e:
                    print(f"Error processing {db_file}: {str(e)}")
                    continue
            
            if row_count > 2:
                title = "تقرير بيانات الطلاب"
                self.style_worksheet(ws, title)
                
                # Adjust column widths
                for idx, column in enumerate(ws.columns, 1):
                    max_length = 0
                    column = list(column)
                    for cell in column:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except:
                            pass
                    adjusted_width = (max_length + 2) * 1.2
                    ws.column_dimensions[get_column_letter(idx)].width = adjusted_width
                
                os.makedirs('exports', exist_ok=True)
                export_path = f"exports/students_{datetime.now().strftime('%Y%m%d')}.xlsx"
                wb.save(export_path)
                messagebox.showinfo("تم", f"تم تصدير بيانات الطلاب إلى {export_path}")
                top = CTkToplevel()
                top.title(f"عرض ملف Excel: {export_path}")
                top.attributes("-topmost", True)
                top.geometry("900x600")
                frame = CTkScrollableFrame(top, fg_color="white")
                frame.pack(fill="both", expand=True, padx=10, pady=10)
                ExcelReader(frame, file_path=export_path)
            else:
                messagebox.showinfo("تنبيه", "لا يوجد طلاب لتصدير بياناتهم")
                
        except Exception as e:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء تصدير البيانات: {str(e)}")

    def update_stats(self):
        total_gains = 0
        total_sessions = 0
        total_students = 0
        
        # Get today's date in the format used in the database
        today = datetime.now()
        today_str = f"{today.day:02d} \\ {today.month:02d} \\ {today.year}"
        
        # Get all .db files in the directory
        db_files = [f for f in os.listdir() if f.endswith('.db') and f != 'Application.db']
        
        for db_file in db_files:
            try:
                conn = connect(db_file)
                cursor = conn.cursor()
                
                # Count gains for today
                cursor.execute("SELECT Qnt FROM Gains WHERE Date LIKE ?", (today_str + '%',))
                gains = cursor.fetchall()
                for gain in gains:
                    try:
                        total_gains += float(gain[0])
                    except (ValueError, TypeError):
                        continue
                
                # Count sessions for today
                cursor.execute("SELECT COUNT(*) FROM Sessions WHERE Date = ?", (today_str,))
                sessions_count = cursor.fetchone()[0]
                total_sessions += sessions_count
                
                # Count total students
                cursor.execute("SELECT COUNT(*) FROM Students")
                students_count = cursor.fetchone()[0]
                total_students += students_count
                
                conn.close()
            except Exception as e:
                print(f"Error processing {db_file}: {str(e)}")
                continue
        
        # Update the labels
        self.gains_label.configure(text=f"E£ {int(total_gains):,}")
        self.sessions_label.configure(text=str(total_sessions))
        self.students_label.configure(text=str(total_students))

    def open_admin(self):
        """Opens the admin page with login"""
        for widget in self.ss.winfo_children():
            widget.destroy()
        Admin.initialize_login(self.ss)

    def open_custom_gains(self):
        """Opens the custom gains management view"""
        for widget in self.ss.winfo_children():
            widget.destroy()
        from Views.CustomGains import CustomGainsView
        CustomGainsView(self.ss).pack(fill=BOTH, expand=True)

    def open_pending_codes(self):
        for widget in self.ss.winfo_children():
            widget.destroy()
        def back_to_team_selector():
            for w in self.ss.winfo_children():
                w.destroy()
            TeamView(self.ss)
        PendingCodesView(self.ss, back_callback=back_to_team_selector)

    def AddNewTeam(self):
            xshow1 = CTkToplevel()
            xshow1.attributes("-topmost", True)
            xshow1.title("إضافة مرحلة جديدة")
            
            Form = CTkScrollableFrame(xshow1, height=500, width=450)
            
            CTkLabel(Form, text="إضافة مرحلة", font=("Cairo Medium", 23)).pack(pady=10)
            
            name = Entries.UpperLabeledEntry(
                Form,
                "اسم المرحلة"
            )
            name.pack(pady=5, padx=14, fill=X)
             
            type_of_payment = StringVar(value="بالحصة")

            def AddTeam():
                db.conn    = connect(f"{name.get_input()}.db")
                db.ref     = db.conn.cursor()
                if type_of_payment.get() == "بالشهر":
                    db.Create("بالشهر")
                else:
                    db.Create()
                
                db.conn = connect(f"Application.db")
                db.ref  = db.conn.cursor()
                db.ref.execute("SELECT * FROM PGroups")
                db.ref.execute("INSERT INTO PGroups VALUES (?,?,?,?)", (len(db.ref.fetchall()) +1 ,name.get_input(), type_of_payment.get(), 0))
                db.conn.commit()
                self.Reload()
                
                xshow1.destroy()
            
            CTkButton(Form, font=("Cairo Medium", 20), text="إضافة المرحلة", command = AddTeam).pack(fill=X, padx=15, pady=(50,0))
            
            Form.pack(padx=10, pady=10, fill=X)
    def Reload(self):
        for x in self.Buttons.winfo_children():
            x.destroy()
        db.conn = connect(f"Application.db")
        db.ref  = db.conn.cursor()
        
        db.ref.execute('SELECT * FROM PGroups')
        Groups = db.ref.fetchall()
        print(Groups)
        for Group in Groups:
            
            def OpenDashboard(database_name):
                for x in self.ss.winfo_children():
                    x.destroy()
                
                change_distnation(database_name)
                print(database_name)
                Dashboard.Dashboard(self.ss).pack(fill=BOTH, expand=True)
            CTkButton(self.Buttons,command= lambda dn = f"{Group[1]}.db" : OpenDashboard(dn), text=Group[1], font=("Tajawal Medium", 40)).pack(pady=10, ipadx=20)
            