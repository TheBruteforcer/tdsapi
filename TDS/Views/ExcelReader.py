from customtkinter import *
from openpyxl import load_workbook
import tkinter.filedialog as fd
import tkinter.messagebox as messagebox
from CTkTable.ctktable import CTkTable
from openpyxl.cell.cell import MergedCell

# Independent Excel Reader View

def ExcelReader(parent: CTkFrame, file_path=None):
    def show_excel(file_path):
        try:
            wb = load_workbook(file_path)
            ws = wb.active
            # Prepare a 2D array for the data
            max_row = ws.max_row
            max_col = ws.max_column
            data = [['' for _ in range(max_col)] for _ in range(max_row)]
            # Fill in merged cells (only top-left gets value, rest are empty)
            for merged_range in ws.merged_cells.ranges:
                min_row, min_col, max_row_m, max_col_m = merged_range.min_row, merged_range.min_col, merged_range.max_row, merged_range.max_col
                value = ws.cell(row=min_row, column=min_col).value
                for r in range(min_row, max_row_m + 1):
                    for c in range(min_col, max_col_m + 1):
                        if r == min_row and c == min_col:
                            data[r-1][c-1] = str(value) if value is not None else ''
                        else:
                            data[r-1][c-1] = ''
            # Fill in the rest of the cells (non-merged or overwrite with actual value)
            for r in range(1, ws.max_row + 1):
                for c in range(1, ws.max_column + 1):
                    cell = ws.cell(row=r, column=c)
                    if isinstance(cell, MergedCell):
                        continue  # Already handled in merged cells loop
                    if not cell.is_date:
                        value = cell.value
                    else:
                        value = cell.value.strftime('%Y-%m-%d %H:%M:%S') if cell.value else ''
                    if value is not None:
                        data[r-1][c-1] = str(value)
            if not data or all(all(cell == '' for cell in row) for row in data):
                messagebox.showinfo("فارغ", "الملف لا يحتوي على بيانات.")
                return
            for widget in parent.winfo_children():
                widget.destroy()
            table = CTkTable(parent, values=data, font=("Cairo Medium", 14), width=120, height=35, corner_radius=0)
            table.pack(fill="both", expand=True, padx=10, pady=10)
        except Exception as e:
            messagebox.showerror("خطأ", f"تعذر قراءة الملف: {e}")
    for widget in parent.winfo_children():
        widget.destroy()
    if file_path:
        show_excel(file_path)
    else:
        def open_and_show():
            file_path = fd.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
            if not file_path:
                return
            show_excel(file_path)
        btn = CTkButton(parent, text="اختر ملف Excel", font=("Cairo Medium", 16), fg_color="#1976D2", command=open_and_show)
        btn.pack(fill=X, padx=20, pady=20) 