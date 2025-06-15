from customtkinter import *
from PIL import Image
from Widgets import Entries
from Proccesors import *
from requests import *
from Proccesors.Database.db import *
from threading import Thread
from Proccesors.Genric.SearchStudent import *
from AlphaControllers.EntryValidators import CheckBlank
from Views import AddStdForm, Spends, AddSpend, Sessions, SessionsSettings, Exams, Money, Groups, Dashboard
from time import localtime
from json import loads, dumps, load
from tkinter import messagebox
import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
import os
from Views.ExcelReader import ExcelReader

def sanitize_group_name(group_name):
    """Sanitizes group name for use in file paths"""
    return group_name.strip().replace("|", "_").replace(" ", "_")

def export_all_students_data():
    try:
        ref.execute("SELECT * FROM Students")
        students = ref.fetchall()
        # Get all exams
        ref.execute("SELECT * FROM Exams")
        exams = ref.fetchall()
        # Get all sessions (attendance)
        ref.execute("SELECT * FROM Sessions")
        sessions = ref.fetchall()
        session_titles = {s[0]: s[1] for s in sessions}
        # Prepare workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "بيانات الطلاب"
        ws.sheet_view.rightToLeft = True
        # Headers
        headers = [
            "ID", "الاسم", "رقم الهاتف", "ولي الأمر", "الفرقة", "سعر الاشتراك", "نوع الدفع", "الدرجات (الامتحانات)", "الحضور (الجلسات)", "المدفوعات"
        ]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True, name='Cairo', size=12)
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.fill = PatternFill(start_color="1976D2", end_color="1976D2", fill_type="solid")
            ws.column_dimensions[chr(64 + col)].width = 25
        # Data
        for idx, student in enumerate(students, 2):
            # Basic info
            ws.cell(row=idx, column=1, value=student[0])  # ID
            ws.cell(row=idx, column=2, value=student[1])  # Name
            ws.cell(row=idx, column=3, value=student[3])  # Phone
            ws.cell(row=idx, column=4, value=student[4])  # Parent phone
            ws.cell(row=idx, column=5, value=student[8])  # Group
            ws.cell(row=idx, column=6, value=student[5])  # Subscription amount
            # Payment type
            payment_type = "بيدفع بالحصة" if student[10] == "BySession" else "شهري/مسبق"
            ws.cell(row=idx, column=7, value=payment_type)
            # Exams (grades)
            student_degrees = []
            for exam in exams:
                try:
                    exam_data = loads(exam[7])
                    for deg in exam_data['degrees']:
                        if student[1] in deg:
                            student_degrees.append(f"{exam[1]}: {deg[student[1]]}/{exam_data['max']}")
                except Exception:
                    continue
            ws.cell(row=idx, column=8, value="; ".join(student_degrees))
            # Attendance
            try:
                attendance = loads(student[9]) if student[9] else []
                attended_sessions = []
                for record in attendance:
                    if isinstance(record, dict):
                        for k, v in record.items():
                            if v == 'Attended.':
                                attended_sessions.append(session_titles.get(k, k))
                ws.cell(row=idx, column=9, value=", ".join(attended_sessions))
            except Exception:
                ws.cell(row=idx, column=9, value="-")
            # Payments
            try:
                if student[10] == "BySession":
                    ws.cell(row=idx, column=10, value="بالحصة")
                else:
                    payments = loads(student[10])
                    payment_strs = []
                    for p in payments:
                        state = p.get('state', '-')
                        month = p.get('month', '-')
                        date_payed = p.get('date_payed', '-')
                        expire = p.get('expire', '-')
                        payment_strs.append(f"{month}: {state} (دفع: {date_payed}, استحقاق: {expire})")
                    ws.cell(row=idx, column=10, value="; ".join(payment_strs))
            except Exception:
                ws.cell(row=idx, column=10, value="-")
        # Save file
        if not os.path.exists('Reports'):
            os.makedirs('Reports')
        filename = f"Reports/تقرير_جميع_الطلاب_{datetime.datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        wb.save(filename)
        messagebox.showinfo("تم", f"تم تصدير بيانات الطلاب بنجاح إلى:\n{filename}")
        top = CTkToplevel()
        top.title(f"عرض ملف Excel: {filename}")
        top.attributes("-topmost", True)
        top.geometry("900x600")
        frame = CTkScrollableFrame(top, fg_color="white")
        frame.pack(fill="both", expand=True, padx=10, pady=10)
        ExcelReader(frame, file_path=filename)
    except Exception as e:
        messagebox.showerror("خطأ", f"حدث خطأ أثناء تصدير البيانات: {str(e)}")

def StudentsPage(parent : CTkScrollableFrame):
    # GUI
    for widget in parent.winfo_children():
        widget.destroy()

    # -- Title Frame -- #
    UpperingFrame = CTkFrame(parent, fg_color='white', corner_radius=20)
    UpperingFrame.pack(fill=X, padx=15, pady=10)

    CTkLabel(UpperingFrame, text='', image=CTkImage(Image.open('TDSAssets/General/student2.png'), size=(70,70))).pack(side=RIGHT, padx=10, pady=10)
    CTkLabel(UpperingFrame, text="صفحة جرد الطلاب", font=('Cairo Medium', 27)).pack(side=RIGHT, padx=10, pady=(15,10))
    def back():
        for x in parent.winfo_children():
            x.destroy()
        Dashboard.Dashboard(parent).pack(fill=BOTH, expand=True)
    CTkButton(UpperingFrame, text='الرجوع الي الرئيسية', font=('Cairo Medium', 16), command=back).pack(side=LEFT, padx=10, ipady=10, ipadx=17)
    # Export button
    CTkButton(UpperingFrame, text='تصدير بيانات الطلاب', font=('Cairo Medium', 16), fg_color='#4CAF50', hover_color='#388E3C', command=export_all_students_data).pack(side=LEFT, padx=10, ipady=10, ipadx=17)

    MainingFrame = CTkFrame(parent, fg_color="white", height=630, corner_radius=20)
    StudentsList = CTkScrollableFrame(MainingFrame, height=530)
    stdimg = CTkImage(Image.open('TDSAssets/General/student2.png'), size=(40,40))
    def View(StudentsData):
        for XCV in StudentsList.winfo_children():
            XCV.destroy()
        for Student in StudentsData:
            paid = False
            mm = str(datetime.datetime.now().month)
            mm.replace("0", "")
            if Student[10] != "BySession":
                for month in loads(Student[10]):
                    if mm == str(month["month"]).replace("شهر", "").strip():
                        paid = True if month["state"] == "تم" else False
            CustomViewFrame = CTkFrame(StudentsList, fg_color="#fffff0")
            CTkLabel(CustomViewFrame, text='', image=stdimg).pack(side=RIGHT, padx=20, pady=20)
            _CL1 = CTkFrame(CustomViewFrame, fg_color="transparent")
            _CL1.pack(fill=Y, padx=5, pady=5, side=RIGHT)
            CTkLabel(_CL1, text=Student[1], font=("Cairo Medium", 20), anchor="e").pack(pady=(15,2), fill=X) 
            CTkLabel(_CL1, text=Student[3], font=("Cairo Medium", 13), anchor="e").pack(fill=X, pady=(3,2)) 
            _CL2 = CTkFrame(CustomViewFrame, fg_color="transparent")
            _CL2.pack(fill=Y, padx=(5,15), pady=5, side=RIGHT)
            CTkLabel(_CL2, text=f"الفرقة : {Student[8]}", font=("Cairo Medium", 16), anchor="e").pack(pady=(15,2), fill=X) 
            if Student[10] != "BySession":
                CTkFrame(CustomViewFrame, fg_color = "green" if paid else "red", corner_radius=180, width=20, height=20).pack(padx=(10, 30), side=RIGHT)
                CTkLabel(CustomViewFrame, text=f"الطالب دافع الشهر" if paid else "الطالب مدفعش الشهر", font=("Cairo Medium", 16)).pack(side=RIGHT) 
            CustomViewFrame.pack(fill=X, pady=(5,5), padx=10)
    # Show all students by default
    ref.execute("SELECT * FROM Students")
    View(ref.fetchall())
    StudentsList.pack(fill=X, padx=10, pady=(0, 10))
    MainingFrame.pack(padx=15, pady=(0, 10), fill=X)
    FooterFrame = CTkFrame(parent , fg_color = "white")
    FLabelingFrame = CTkFrame(FooterFrame, fg_color="#FFFFA0", corner_radius=10)
    CTkLabel(FLabelingFrame, text="الأدوات المتاحة للجرد", font=("Cairo Medium", 21)).pack(pady=5)
    FLabelingFrame.pack(fill=X, padx=10, pady=10)
    DeleteNAStudentsFrame = CTkFrame(FooterFrame, fg_color="#F2F2F2")
    _DNAF1 = CTkFrame(DeleteNAStudentsFrame, fg_color="transparent")
    _DNAF1.pack(side= RIGHT, pady=10, padx=10)
    CTkLabel(_DNAF1, text=f"حذف جميع الطلاب الغير نشطين", font=("Cairo Medium", 25), anchor="e").pack(pady=(15,2), fill=X) 
    CTkLabel(_DNAF1, text=f"سيتم حذف الطلاب الذين يحملون نسبة حضور أقل من المتوسطة . هذه الخطوة لا رجعة فيها", font=("Cairo Medium", 14), anchor="e").pack(pady=(3,2), fill=X) 
    def DeleteNAStudents():
        messagebox.showinfo("غير متاح", "تم تعطيل هذه الخاصية مؤقتًا")
        return
    CTkButton(DeleteNAStudentsFrame,command=DeleteNAStudents, text="بدء العملية", fg_color="red", hover_color="red", text_color="white", font=("Cairo Medium", 16)).pack(side=LEFT, padx=20, pady=10, ipady=20)
    DeleteNAStudentsFrame.pack(fill=X, padx=10, pady=10)
    FooterFrame.pack(fill=X, padx=15)





